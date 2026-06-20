"""Streamlit UI for the Client Onboarding & KYC Agent."""

from __future__ import annotations

import json
import logging
import os
import subprocess
import sys
from io import BytesIO
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parent
VENV_PYTHON = PROJECT_ROOT / ".venv" / "Scripts" / "python.exe"


def _bootstrap_project_runner() -> None:
    if os.environ.get("STREAMLIT_SERVER_PORT"):
        return
    if sys.executable.lower() == str(VENV_PYTHON).lower():
        return
    if not VENV_PYTHON.exists():
        return
    if "python.exe" not in sys.executable.lower() and "pythonw.exe" not in sys.executable.lower():
        return
    subprocess.run(
        [
            str(VENV_PYTHON),
            "-m",
            "streamlit",
            "run",
            str(PROJECT_ROOT / "app.py"),
            "--server.headless",
            "true",
            "--browser.gatherUsageStats",
            "false",
        ],
        check=False,
    )
    raise SystemExit(0)


_bootstrap_project_runner()

import pandas as pd
import streamlit as st
from pydantic import ValidationError

from kyc_agents.onboarding_agent import LIVE_AGENTS_AVAILABLE, run_onboarding_agent
from schemas.client_profile import ClientProfile, ProcessingResult
from tools.followup_generator import generate_followup_questions
from tools.gap_analyzer import analyze_profile_gaps

try:
    from dotenv import load_dotenv
except Exception:  # pragma: no cover - optional dependency
    load_dotenv = None

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover - optional dependency
    PdfReader = None

try:
    from docx import Document
except Exception:  # pragma: no cover - optional dependency
    Document = None


logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s - %(message)s")
logger = logging.getLogger(__name__)

if load_dotenv:
    load_dotenv()

st.set_page_config(page_title="Client Onboarding & KYC Agent", page_icon="WA", layout="wide")


SAMPLE_DIR = Path(__file__).parent / "sample_data"
AUDIO_SUFFIXES = {".mp3", ".mp4", ".mpeg", ".mpga", ".m4a", ".wav", ".webm"}
SUPPORTED_UPLOAD_TYPES = ["txt", "pdf", "docx", "xlsx", "xls", "csv", "mp3", "mp4", "mpeg", "mpga", "m4a", "wav", "webm"]


def extract_text_from_upload(uploaded_file: Any) -> tuple[str, list[str], dict[str, str]]:
    """Extract text from supported uploads and gracefully report unsupported files."""
    if uploaded_file is None:
        return "", [], {}

    suffix = Path(uploaded_file.name).suffix.lower()
    warnings: list[str] = []
    details = {"name": uploaded_file.name, "type": suffix or "unknown", "status": "Ready"}

    if suffix == ".txt":
        text = uploaded_file.getvalue().decode("utf-8", errors="replace")
        details["status"] = f"{len(text):,} characters extracted"
        return text, warnings, details

    if suffix == ".pdf":
        if PdfReader is None:
            details["status"] = "PDF parser unavailable"
            return "", ["PDF support requires the pypdf package. Install project requirements and retry."], details
        try:
            reader = PdfReader(BytesIO(uploaded_file.getvalue()))
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()
            if not text:
                warnings.append("No selectable text was found in the PDF. Scanned PDFs may need OCR.")
            details["status"] = f"{len(reader.pages)} pages, {len(text):,} characters extracted"
            return text, warnings, details
        except Exception as exc:
            logger.exception("PDF parsing failed")
            details["status"] = "PDF parse failed"
            return "", [f"Could not read the PDF: {exc}"], details

    if suffix == ".docx":
        if Document is None:
            details["status"] = "Word parser unavailable"
            return "", ["Word upload support requires python-docx. Install project requirements and retry."], details
        try:
            document = Document(BytesIO(uploaded_file.getvalue()))
            paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        paragraphs.append(" | ".join(cells))
            text = "\n".join(paragraphs)
            details["status"] = f"{len(paragraphs)} Word blocks extracted"
            return text, warnings, details
        except Exception as exc:
            logger.exception("Word parsing failed")
            details["status"] = "Word parse failed"
            return "", [f"Could not read the Word document: {exc}"], details

    if suffix in {".xlsx", ".xls"}:
        try:
            sheets = pd.read_excel(BytesIO(uploaded_file.getvalue()), sheet_name=None, dtype=str)
            chunks: list[str] = []
            for sheet_name, frame in sheets.items():
                frame = frame.dropna(how="all").fillna("")
                chunks.append(f"Sheet: {sheet_name}")
                chunks.append(frame.to_csv(index=False))
            text = "\n".join(chunks).strip()
            details["status"] = f"{len(sheets)} workbook sheets extracted"
            return text, warnings, details
        except Exception as exc:
            logger.exception("Excel parsing failed")
            details["status"] = "Excel parse failed"
            return "", [f"Could not read the Excel workbook: {exc}"], details

    if suffix == ".csv":
        try:
            frame = pd.read_csv(BytesIO(uploaded_file.getvalue()), dtype=str).fillna("")
            text = frame.to_csv(index=False)
            details["status"] = f"{len(frame):,} CSV rows extracted"
            return text, warnings, details
        except Exception as exc:
            logger.exception("CSV parsing failed")
            details["status"] = "CSV parse failed"
            return "", [f"Could not read the CSV file: {exc}"], details

    if suffix in AUDIO_SUFFIXES:
        text, audio_warnings = transcribe_audio_upload(uploaded_file)
        details["status"] = f"{len(text):,} transcript characters" if text else "Transcription unavailable"
        return text, audio_warnings, details

    details["status"] = "Unsupported"
    return "", [f"Unsupported file type: {suffix or 'unknown'}. Please upload TXT, PDF, Word, Excel, CSV, or audio."], details


def extract_text_from_uploads(uploaded_files: list[Any]) -> tuple[str, list[str], list[dict[str, str]]]:
    """Extract and combine text from a multi-file upload queue."""
    combined: list[str] = []
    warnings: list[str] = []
    manifest: list[dict[str, str]] = []
    for uploaded_file in uploaded_files:
        text, file_warnings, details = extract_text_from_upload(uploaded_file)
        manifest.append(details)
        warnings.extend(file_warnings)
        if text.strip():
            combined.append(f"--- Source: {uploaded_file.name} ---\n{text.strip()}")
    return "\n\n".join(combined), warnings, manifest


def transcribe_audio_upload(uploaded_file: Any) -> tuple[str, list[str]]:
    """Transcribe audio when the optional OpenAI package and API key are available."""
    if not os.getenv("OPENAI_API_KEY"):
        return "", ["Audio upload received, but OPENAI_API_KEY is not set. Paste a transcript or enable live transcription."]
    try:
        from openai import OpenAI
    except Exception:
        return "", ["Audio transcription requires the openai package. Install project requirements and retry."]

    try:
        client = OpenAI()
        transcript = client.audio.transcriptions.create(
            model=os.getenv("OPENAI_TRANSCRIPTION_MODEL", "gpt-4o-mini-transcribe"),
            file=(uploaded_file.name, BytesIO(uploaded_file.getvalue())),
        )
        return getattr(transcript, "text", ""), []
    except Exception as exc:
        logger.exception("Audio transcription failed")
        return "", [f"Audio transcription failed. Paste a transcript instead. Details: {exc}"]


def load_sample_text(filename: str) -> str:
    return (SAMPLE_DIR / filename).read_text(encoding="utf-8")


def render_confidence(profile: ClientProfile, field_name: str) -> str:
    score = profile.confidence_scores.get(field_name)
    if score is None:
        return "Not extracted"
    return f"{score:.0%}"


def render_workflow_bar(has_input: bool, has_result: bool, approved: bool) -> None:
    steps = [
        ("1", "Collect", "Upload files or select a scenario", has_input),
        ("2", "Analyze", "Extract profile and KYC gaps", has_result),
        ("3", "Review", "Advisor edits and follow-up", has_result),
        ("4", "Export", "Download after approval", approved),
    ]
    cols = st.columns(4)
    for col, (number, label, caption, active) in zip(cols, steps):
        with col:
            st.markdown(f"### {'Complete' if active else number}. {label}")
            st.caption(caption)


def profile_to_report(result: ProcessingResult) -> str:
    profile = result.profile
    lines = [
        "# Client Onboarding & KYC Review",
        "",
        "Advisor approval is mandatory. This report is a first-pass onboarding analysis, not a KYC approval or suitability decision.",
        "",
    ]
    for section, body in result.advisor_summary.items():
        lines.extend([f"## {section}", body, ""])
    lines.extend(["## Structured Profile", ""])
    for field_name, value in profile.model_dump().items():
        if field_name in {"confidence_scores", "source_attribution"}:
            continue
        lines.append(f"- {field_name.replace('_', ' ').title()}: {value}")
    lines.extend(["", "## Source Attribution", ""])
    for field_name, source in profile.source_attribution.items():
        lines.append(f"- {field_name.replace('_', ' ').title()}: {source}")
    return "\n".join(lines)


def apply_advisor_edits(profile: ClientProfile) -> ClientProfile:
    with st.form("advisor_edits"):
        st.subheader("Advisor Edits")
        risk_tolerance = st.selectbox(
            "Risk Tolerance",
            ["", "Conservative", "Moderate", "Balanced", "Aggressive"],
            index=["", "Conservative", "Moderate", "Balanced", "Aggressive"].index(profile.risk_tolerance)
            if profile.risk_tolerance in {"Conservative", "Moderate", "Balanced", "Aggressive"}
            else 0,
        )
        goals = st.text_area("Goals", value=", ".join(profile.financial_goals), height=88)
        dependents = st.text_input("Dependents", value=profile.dependents or "")
        horizon = st.text_input("Horizon", value=profile.investment_horizon or "")
        liquidity = st.text_input("Liquidity Needs", value=profile.liquidity_needs or "")
        approved_for_export = st.checkbox("I reviewed these extracted fields and approve exporting this profile.")
        submitted = st.form_submit_button("Update Profile")

    if submitted:
        data = profile.model_dump()
        data.update(
            {
                "risk_tolerance": risk_tolerance or None,
                "financial_goals": [item.strip() for item in goals.split(",") if item.strip()],
                "dependents": dependents or None,
                "investment_horizon": horizon or None,
                "liquidity_needs": liquidity or None,
            }
        )
        st.session_state.profile_approved = approved_for_export
        try:
            return ClientProfile.model_validate(data)
        except ValidationError as exc:
            st.error(f"Could not apply edits: {exc}")
    return profile


def render_result(result: ProcessingResult) -> None:
    profile = result.profile
    if result.warnings:
        for warning in result.warnings:
            st.warning(warning)

    metric_cols = st.columns(4)
    metric_cols[0].metric("Missing Fields", len(result.gap_analysis.missing_fields))
    metric_cols[1].metric("Follow-Ups", len(profile.followup_questions))
    metric_cols[2].metric("Extracted Fields", sum(1 for key in ["client_name", "age", "occupation", "annual_income", "risk_tolerance", "investment_horizon", "liquidity_needs"] if profile.value_for_field(key) not in (None, "", [])))
    metric_cols[3].metric("Mode", result.mode)

    tabs = st.tabs(["Review", "Advisor Edits", "Agent Logs", "Evidence", "Export"])

    with tabs[0]:
        top_cols = st.columns([1.2, 0.8])
        with top_cols[0]:
            st.subheader("Advisor Summary")
            for title in ["Client Overview", "Goals", "Risk Profile"]:
                st.markdown(f"**{title}**")
                st.write(result.advisor_summary.get(title, "No information available."))
        with top_cols[1]:
            st.subheader("Missing Information")
            if result.gap_analysis.missing_fields:
                for field in result.gap_analysis.missing_fields:
                    st.warning(field)
            else:
                st.success("No core KYC gaps detected. Advisor approval is still required.")
            for area in result.gap_analysis.risk_areas:
                st.info(area)

        st.subheader("Follow-Up Questions")
        if profile.followup_questions:
            for index, question in enumerate(profile.followup_questions, start=1):
                st.write(f"{index}. {question}")
            st.code("\n".join(profile.followup_questions), language="text")
        else:
            st.success("No follow-up questions generated.")

    with tabs[1]:
        edited_profile = apply_advisor_edits(profile)
        if edited_profile != profile:
            gap_analysis = analyze_profile_gaps(edited_profile)
            edited_profile.missing_information = gap_analysis.missing_fields
            edited_profile.followup_questions = generate_followup_questions(gap_analysis.missing_fields)
            updated_result = ProcessingResult(
                profile=edited_profile,
                gap_analysis=gap_analysis,
                advisor_summary={
                    "Client Overview": result.advisor_summary.get("Client Overview", ""),
                    "Goals": f"Financial Goals: {', '.join(edited_profile.financial_goals) or 'No reliable information extracted.'}\nInvestment Horizon: {edited_profile.investment_horizon or 'No reliable information extracted.'}",
                    "Risk Profile": f"Risk Tolerance: {edited_profile.risk_tolerance or 'No reliable information extracted.'}\nLiquidity Needs: {edited_profile.liquidity_needs or 'No reliable information extracted.'}",
                    "Missing Information": "\n".join(f"- {field}" for field in gap_analysis.missing_fields) or "No core KYC gaps detected. Advisor approval remains mandatory.",
                    "Recommended Follow-Up": "\n".join(f"- {question}" for question in edited_profile.followup_questions) or "No follow-up questions generated.",
                },
                mode=result.mode,
                warnings=result.warnings,
            )
            st.session_state.result = updated_result
            st.rerun()

    with tabs[2]:
        st.subheader("Live Agent Orchestration Log")
        if result.run_log:
            for entry in result.run_log:
                st.code(entry, language="text")
        else:
            st.info("No agent log entries were captured for this run.")
        if result.last_agent:
            st.metric("Last Agent", result.last_agent)
        if result.trace_hint:
            st.caption(result.trace_hint)

    with tabs[3]:
        st.subheader("Confidence & Sources")
        source_rows = []
        for field_name in ["client_name", "annual_income", "risk_tolerance", "investment_horizon", "liquidity_needs", "financial_goals", "existing_investments"]:
            source_rows.append(
                {
                    "Field": field_name.replace("_", " ").title(),
                    "Confidence": render_confidence(profile, field_name),
                    "Source": profile.source_attribution.get(field_name, "No source captured"),
                }
            )
        st.dataframe(source_rows, hide_index=True, width="stretch")
        with st.expander("Structured JSON Output", expanded=False):
            st.json(profile.model_dump())

    with tabs[4]:
        st.subheader("Export")
        approved = st.session_state.get("profile_approved", False)
        export_payload = json.dumps(profile.model_dump(), indent=2)
        report_payload = profile_to_report(result)
        export_cols = st.columns(2)
        with export_cols[0]:
            st.download_button(
                "Download JSON",
                data=export_payload,
                file_name="client_profile.json",
                mime="application/json",
                disabled=not approved,
            )
        with export_cols[1]:
            st.download_button(
                "Download Report",
                data=report_payload,
                file_name="client_onboarding_report.md",
                mime="text/markdown",
                disabled=not approved,
            )
        if not approved:
            st.caption("Exports unlock after advisor review approval in the edit form.")


def main() -> None:
    st.title("Client Onboarding & KYC Agent")
    st.write("A guided review workbench for turning onboarding files into an advisor-approved KYC profile.")

    with st.sidebar:
        st.header("Workspace")
        scenario_files = {
            "None": "",
            "Retirement-focused investor": "scenario_1_retirement_complete.txt",
            "Incomplete onboarding": "scenario_2_incomplete.txt",
            "High-net-worth aggressive investor": "scenario_3_hnw_aggressive.txt",
        }
        selected_sample = st.selectbox("Synthetic Scenario", list(scenario_files))
        st.caption("Live mode requires openai-agents and OPENAI_API_KEY.")
        api_key_ready = bool(os.getenv("OPENAI_API_KEY"))
        st.markdown("### Live Orchestration")
        if LIVE_AGENTS_AVAILABLE and api_key_ready:
            st.success("Ready")
            st.caption("The app will stream agent updates, tool calls, and handoffs in the Agent Logs tab.")
        else:
            st.warning("Not ready")
            if not LIVE_AGENTS_AVAILABLE:
                st.caption("Activate the project venv so the openai-agents package is available.")
            if not api_key_ready:
                st.caption("Set OPENAI_API_KEY in your .env file to enable live agent runs.")

    sample_text = load_sample_text(scenario_files[selected_sample]) if scenario_files[selected_sample] else ""
    uploaded_files = st.file_uploader(
        "Upload onboarding dataset",
        type=SUPPORTED_UPLOAD_TYPES,
        accept_multiple_files=True,
        help="Supported: TXT, PDF, Word DOCX, Excel XLSX/XLS, CSV, and audio files.",
    )
    uploaded_text, upload_warnings, upload_manifest = extract_text_from_uploads(uploaded_files or [])

    render_workflow_bar(bool(uploaded_text or sample_text), "result" in st.session_state, st.session_state.get("profile_approved", False))

    if upload_manifest:
        st.subheader("File Queue")
        st.dataframe(upload_manifest, hide_index=True, width="stretch")

    default_text = uploaded_text or sample_text
    raw_text = st.text_area("Onboarding text assembled from uploads and notes", value=default_text, height=260)

    action_cols = st.columns([0.25, 0.25, 0.5])
    process = action_cols[0].button("Process Onboarding", type="primary", width="stretch")
    if action_cols[1].button("Clear Review", width="stretch"):
        st.session_state.pop("result", None)
        st.session_state.profile_approved = False
        st.rerun()
    if process:
        st.session_state.profile_approved = False
        with st.spinner("Analyzing onboarding information..."):
            result = run_onboarding_agent(raw_text)
            result.warnings.extend(upload_warnings)
            st.session_state.result = result

    if "result" in st.session_state:
        render_result(st.session_state.result)


if __name__ == "__main__":
    main()
