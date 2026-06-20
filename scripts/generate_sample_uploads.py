"""Generate sample onboarding upload files for demos."""

from __future__ import annotations

import csv
import subprocess
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "sample_uploads"
OUT.mkdir(exist_ok=True)

PROFILE_TEXT = """Client Name: Aisha Patel
Age: 51
Occupation: Orthopedic surgeon
Marital Status: Married
Dependents: 2 children
Annual Income: $875,000 annually
Risk Tolerance: Balanced
Investment Horizon: 12 years
Liquidity Needs: Keep $350,000 liquid for taxes, tuition, and emergency reserves.

Aisha wants retirement planning, education funding, tax efficiency, and long-term growth.
Existing investments include a 401(k), IRA, taxable brokerage account, ETFs, bonds, and real estate.
"""


def write_pdf() -> None:
    pdf_path = OUT / "aisha_patel_onboarding_notes.pdf"
    doc = canvas.Canvas(str(pdf_path), pagesize=letter)
    width, height = letter
    text = doc.beginText(72, height - 72)
    text.setFont("Helvetica", 11)
    for line in PROFILE_TEXT.splitlines():
        text.textLine(line)
    doc.drawText(text)
    doc.save()


def write_docx() -> None:
    doc_path = OUT / "aisha_patel_meeting_transcript.docx"
    document = Document()
    document.add_heading("Advisor Meeting Transcript", level=1)
    document.add_paragraph(PROFILE_TEXT)
    document.add_paragraph("Advisor note: confirm concentrated employer stock exposure before final recommendation.")
    document.save(doc_path)


def write_xlsx() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Client Intake"
    rows = [
        ("Field", "Value"),
        ("Client Name", "Aisha Patel"),
        ("Age", "51"),
        ("Occupation", "Orthopedic surgeon"),
        ("Annual Income", "$875,000"),
        ("Risk Tolerance", "Balanced"),
        ("Investment Horizon", "12 years"),
        ("Liquidity Needs", "$350,000 reserve"),
        ("Existing Investments", "401(k), IRA, brokerage, ETFs, bonds, real estate"),
    ]
    for row in rows:
        sheet.append(row)
    workbook.save(OUT / "aisha_patel_intake_workbook.xlsx")


def write_csv() -> None:
    with (OUT / "aisha_patel_holdings.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Account", "Asset Type", "Notes"])
        writer.writerow(["401(k)", "Retirement", "Employer plan with index funds"])
        writer.writerow(["Taxable brokerage", "ETFs", "Long-term growth allocation"])
        writer.writerow(["Real estate", "Property", "Rental property with moderate liquidity"])


def write_audio() -> None:
    wav_path = OUT / "aisha_patel_voice_note.wav"
    script = f"""
Add-Type -AssemblyName System.Speech
$speaker = New-Object System.Speech.Synthesis.SpeechSynthesizer
$speaker.SetOutputToWaveFile('{wav_path}')
$speaker.Speak('Client Name Aisha Patel. Age 51. Occupation orthopedic surgeon. Risk tolerance balanced. Investment horizon 12 years. Liquidity needs 350 thousand dollars for taxes tuition and emergency reserves.')
$speaker.Dispose()
"""
    try:
        subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True, capture_output=True, text=True)
    except Exception:
        wav_path.write_bytes(b"")


def main() -> None:
    write_pdf()
    write_docx()
    write_xlsx()
    write_csv()
    write_audio()


if __name__ == "__main__":
    main()
